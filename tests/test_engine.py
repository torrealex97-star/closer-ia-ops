import asyncio
import io
import json
import sqlite3
from datetime import datetime, timedelta, timezone
import pytest
from pydantic import SecretStr
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from career import connectors
from career.connectors import AtsCompany, Candidate, Himalayas, WeWorkRemotely, ProviderError, REGISTRY, watchlist_connectors
from career.db import connect, load_run, now
from career.main import app
from career.models import SearchQuery
from career.pipeline import ingest, match, normalize, plain, safe_url
from career.worker import execute
from career.presentation import professional_role


def test_vocabulary_families_recognize_equivalent_titles():
    """Taxonomía WDC 2026-09-23: titulos que no dicen closer pero son el mismo perfil."""
    assert professional_role('Asesor de Admisiones', 'captacion de alumnos bootcamp') == 'closer'
    assert professional_role('Admissions Advisor', 'student recruitment') == 'closer'
    assert professional_role('Enrollment Specialist', 'education sales') == 'closer'
    assert professional_role('Patient Coordinator', 'clinica dental') == 'closer'
    assert professional_role('Treatment Coordinator', 'clinica capilar') == 'closer'
    assert professional_role('Sales Advisor', 'formacion digital') == 'account_executive'
    assert professional_role('Consultor Comercial', 'negociacion comercial') == 'account_executive'
    assert professional_role('Teleoperador Comercial', 'citas agendadas') == 'setter'
    assert professional_role('Growth Manager', 'analitica y producto') is None


def test_match_quality_signals_and_noise():
    base = {'title': 'Closer de ventas', 'company': 'Academia X', 'description': '<p>Remoto.</p>',
            'source_url': 'https://example.com/x', 'locations': ['España'], 'published_at': now(),
            'work_mode': 'remote', 'salary_min': None, 'salary_max': None,
            'salary_currency': None, 'salary_period': None}
    plain_offer = dict(base)
    score, reasons, _ = match(plain_offer, SearchQuery())
    strong = dict(base, description='<p>Remoto. Leads cualificados, fijo más comisiones y contrato indefinido. CRM propio.</p>')
    score2, reasons2, _ = match(strong, SearchQuery())
    assert score2 > score, (reasons, reasons2)
    assert any('proceso comercial' in r.lower() for r in reasons2)
    # Ruido: curso para aprender closing y MLM nunca son oportunidades.
    course = dict(base, title='Curso de Closer', description='<p>Aprende a cerrar ventas.</p>')
    assert match(course, SearchQuery()) is None
    mlm = dict(base, title='Closer de ventas', description='<p>Red multinivel, network marketing.</p>')
    assert match(mlm, SearchQuery()) is None
    scam = dict(base, title='Closer de ventas', description='<p>Inversiones garantizadas.</p>')
    assert match(scam, SearchQuery()) is None


def test_match_recognizes_admissions_and_sector_terms():
    offer = {'title': 'Asesor de Admisiones', 'company': 'Universidad Online',
             'description': '<p>Captación de alumnos, matriculacion. Bootcamp de IA. Remoto.</p>',
             'source_url': 'https://example.com/a', 'locations': ['España'], 'published_at': now(),
             'work_mode': 'remote', 'salary_min': None, 'salary_max': None,
             'salary_currency': None, 'salary_period': None}
    result = match(offer, SearchQuery())
    assert result, 'las ofertas de admisiones deben contar para digital_sales'
    score, reasons, _ = result
    assert score >= 60


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("CAREER_DB", str(tmp_path / "test.sqlite3"))
    monkeypatch.setenv("CAREER_LOCAL_TOKEN", "test-token")
    monkeypatch.setenv("CAREER_ORG_ID", "org-a")
    from career import collection
    async def quote(value, token):
        return 1.30, 0
    monkeypatch.setattr(collection, 'budget_quote', quote)
    with TestClient(app, headers={"Authorization": "Bearer test-token"}) as c:
        yield c


@pytest.fixture
def candidate():
    return Candidate("https://himalayas.app/companies/example/jobs/closer", {
        "guid": "https://himalayas.app/companies/example/jobs/closer", "title": "Sales Closer",
        "companyName": "Example Company", "description": "<p>Spanish-speaking closer. Base plus commission.</p>",
        "locationRestrictions": ["Spain"], "pubDate": now(), "minSalary": 36000,
        "maxSalary": 48000, "currency": "EUR", "salaryPeriod": "annual",
        "applicationLink": "https://example.com/apply",
    })


def create(client, query=None, key="test-request-001"):
    response = client.post("/v1/search-runs", json={"query": query or {}, "idempotency_key": key})
    assert response.status_code == 202, response.text
    return response.json()


def test_auth_and_empty_state(client):
    assert client.get("/v1/config", headers={"Authorization": ""}).status_code == 401
    assert client.get("/v1/search-runs").json() == []
    assert client.get("/v1/config").json()["mode"] == "local"
    # 22 públicas anteriores + FreeHire (capa FREE de SPEC-002, disponible sin configurar).
    assert sum(s["available"] for s in client.get("/v1/config").json()["sources"]) == 23
    freehire = next(s for s in client.get("/v1/config").json()["sources"] if s["code"] == "freehire")
    assert freehire["available"] is True and freehire["requires"] == "none"


def test_profile_terms_feed_apify_inputs(client):
    """El vocabulario canónico alimenta los inputs de Apify del asistente."""
    data = client.get("/v1/profile-terms").json()
    roles = data["roles"]
    assert set(roles) >= {"digital_sales", "closer", "setter", "account_executive"}
    ds = roles["digital_sales"]["terms"]
    # La 'lista mínima' del usuario va primero y cubre los equivalentes clave.
    for term in ["closer", "asesor de admisiones", "admissions advisor", "enrollment advisor",
                 "patient coordinator", "sales advisor", "treatment coordinator"]:
        assert term in ds, term
    assert ds.index("closer") < ds.index("patient coordinator"), "los preferidos primero"
    assert roles["digital_sales"]["label"] == "Closers, setters y comercial online"
    # Los términos del asistente y los del KEYWORDS de integraciones comparten vocabulario.
    from career.integrations import KEYWORDS
    assert KEYWORDS[0] == "closer" and "high ticket closer" in KEYWORDS


def test_google_dorks_presets_ready_for_serpapi(client, monkeypatch):
    # Los presets WDC pertenecen a local-wdc; el fixture client usa org-a.
    monkeypatch.setenv("CAREER_ORG_ID", "local-wdc")
    """Dorks Google (site:linkedin.com/jobs/view con comillas) como presets para SerpApi."""
    from career.source_strategy import source_presets
    presets = [p for p in source_presets()['searches'] if p['source'] == 'google']
    assert len(presets) >= 12, 'la taxonomia debe tener sus dorks preparados'
    for p in presets:
        assert p['dork'].startswith('site:linkedin.com/jobs/view'), p['name']
        assert '"' in p['dork'], 'los dorks usan comillas exactas'
    # el dork de leads cualificados combina cargo + senal de calidad (seccion 11)
    assert any('"leads cualificados"' in p['dork'] for p in presets)
    config = client.get("/v1/config").json()["sources"]
    google = next(s for s in config if s["code"] == "google")
    assert google["available"] is False, "SerpApi aun no esta conectado"
    assert "API key" in google["validation_status"] or "key" in google["validation_status"].lower()
    assert len(google["suggested_searches"]) >= 12
    # el asistente debe poder abrir la fuente con su primer dork como URL
    first = google["suggested_searches"][0]
    assert first["url"].startswith("https://www.google.com/search?q=")


def test_idempotency_and_conflicting_replay(client):
    first = create(client)
    assert create(client)["id"] == first["id"]
    response = client.post("/v1/search-runs", json={"query": {"role": "setter"}, "idempotency_key": "test-request-001"})
    assert response.status_code == 409


def test_raw_is_immutable_and_retry_does_not_duplicate(client, candidate):
    run = create(client)
    for _ in range(2):
        assert ingest(candidate, Himalayas(), SearchQuery(), run["id"], "org-a")
    with connect() as db:
        for table in ["raw_items", "job_postings", "jobs", "search_results"]:
            assert db.execute(f"SELECT count(*) FROM {table}").fetchone()[0] == 1
        with pytest.raises(sqlite3.IntegrityError):
            db.execute("UPDATE raw_items SET payload='{}'")
        with pytest.raises(sqlite3.IntegrityError):
            db.execute("DELETE FROM raw_items")
    assert len(load_run(run["id"], "org-a")["jobs"]) == 1


def test_cross_org_reads_writes_exports_are_denied(client, candidate, monkeypatch):
    run = create(client)
    ingest(candidate, Himalayas(), SearchQuery(), run["id"], "org-a")
    job = load_run(run["id"], "org-a")["jobs"][0]
    monkeypatch.setenv("CAREER_ORG_ID", "org-b")
    assert client.get(f'/v1/search-runs/{run["id"]}').status_code == 404
    assert client.patch(f'/v1/jobs/{job["id"]}/star', json={"starred": True}).status_code == 404
    assert client.post("/v1/exports", json={"run_id": run["id"], "job_ids": [job["id"]], "format": "json"}).status_code == 404
    assert client.get("/v1/search-runs").json() == []
    with connect() as db:
        with pytest.raises(sqlite3.IntegrityError):
            db.execute("INSERT INTO search_results VALUES(?,?,?,?)", ("org-b", run["id"], job["id"], "{}"))


def test_strict_filters_do_not_invent_eligibility(candidate):
    item = normalize("himalayas", candidate.payload)
    assert match(item, SearchQuery(role="closer", market="ES", salary_min=2000, salary_mode="strict"))
    assert match(item, SearchQuery(market="US")) is None
    assert match(item, SearchQuery(salary_min=4000, salary_mode="strict")) is None
    item["salary_currency"] = "USD"
    assert match(item, SearchQuery(salary_min=2000, salary_mode="strict")) is None
    item["locations"] = []
    assert match(item, SearchQuery(market="ES", include_unknown=False)) is None
    assert match(item, SearchQuery(market="ES", include_unknown=True))[2]
    item["published_at"] = (datetime.now(timezone.utc) - timedelta(days=100)).isoformat()
    assert match(item, SearchQuery()) is None


def test_expired_and_misleading_sdr_substring(candidate):
    item = normalize("himalayas", candidate.payload)
    item["expires_at"] = (datetime.now(timezone.utc) - timedelta(days=1)).isoformat()
    assert match(item, SearchQuery()) is None
    item["expires_at"] = None
    item["title"] = "Cloud specialist"
    item["description"] = "AWS SDK and sdrandom"
    assert match(item, SearchQuery(role="sdr")) is None


def test_safe_rendering():
    assert safe_url("javascript:alert(1)") == ""
    assert safe_url("https://user:secret@example.com") == ""
    assert plain("<script>bad()</script><p>Good &amp; safe</p>") == "Good & safe"


def test_sales_team_mention_does_not_turn_engineer_into_sales_job(candidate):
    item = normalize("himalayas", candidate.payload)
    item["title"] = "Software Engineer"
    item["description"] = "Work with sales and account executives. Take a closer look at our company."
    assert match(item, SearchQuery(role="sales")) is None
    item["description"] = "We are hiring a high ticket closer to join our sales team."
    assert match(item, SearchQuery(role="closer"))


def test_partial_failure_and_progressive_results(client, candidate):
    run = create(client)
    class Good(Himalayas):
        async def batches(self, query):
            yield [candidate]
    class Bad(WeWorkRemotely):
        async def batches(self, query):
            await asyncio.sleep(.02)
            assert len(load_run(run["id"], "org-a")["jobs"]) == 1
            raise ProviderError("Prueba: fuente caída")
            yield []
    asyncio.run(execute(run["id"], "org-a", SearchQuery(), {"himalayas": Good(), "weworkremotely": Bad()}))
    result = load_run(run["id"], "org-a")
    assert result["status"] == "partial"
    assert len(result["jobs"]) == 1
    assert result["sources"][1]["status"] == "failed"


def test_exports_selected_only_and_formula_safe(client, candidate):
    run = create(client)
    candidate.payload["companyName"] = "=HYPERLINK(\"https://example.com\")"
    ingest(candidate, Himalayas(), SearchQuery(), run["id"], "org-a")
    job = load_run(run["id"], "org-a")["jobs"][0]
    body = {"run_id": run["id"], "job_ids": [job["id"]]}
    data = client.post("/v1/exports", json={**body, "format": "json"}).json()
    assert len(data) == 1 and data[0]["source_url"] == candidate.external_id
    csv = client.post("/v1/exports", json={**body, "format": "csv"})
    assert "'=HYPERLINK" in csv.text
    xlsx = client.post("/v1/exports", json={**body, "format": "xlsx"})
    workbook = load_workbook(io.BytesIO(xlsx.content))
    assert workbook.active.max_row == 2
    assert workbook.active["B2"].data_type == "s"
    assert client.post("/v1/exports", json={**body, "job_ids": ["nonexistent"], "format": "csv"}).status_code == 422


def test_saved_search_and_cancellation_persist(client):
    saved = client.post("/v1/saved-searches", json={"name": "Closers ES", "query": {"role": "closer", "market": "ES"}})
    assert saved.status_code == 200
    assert client.get("/v1/saved-searches").json()[0]["query"]["market"] == "ES"
    run = create(client)
    assert client.post(f'/v1/search-runs/{run["id"]}/cancel').json()["status"] == "cancelled"
    assert load_run(run["id"], "org-a")["status"] == "cancelled"


def test_expansion_connectors_jobicy_rss(monkeypatch):
    async def fake_fetch(url, params=None, ttl=900):
        if "jobicy.com" in url:
            assert params == {"count": "50", "industry": "seller"}
            return json.dumps({"jobs": [{"id": 1, "url": "https://jobicy.com/jobs/1-closer", "jobTitle": "Closer",
                "companyName": "Acme", "jobDescription": "<p>Buscamos closer</p>", "jobGeo": "Spain",
                "pubDate": "2026-09-22T00:00:00+00:00"}]})
        if "remotefirstjobs" in url:
            return '<rss><channel><item><title>Closer at Acme</title><link>https://remotefirstjobs.com/x</link><description>Buscamos closer</description></item></channel></rss>'
        if "jobspresso" in url:
            return '<rss><channel><item><title>Closer de ventas</title><link>https://jobspresso.co/x</link><description>Remote sales role</description></item></channel></rss>'
        raise AssertionError(url)
    monkeypatch.setattr("career.connectors.fetch", fake_fetch)
    async def collect(connector):
        return [row async for batch in connector.batches(SearchQuery()) for row in batch]
    item = normalize("jobicy", asyncio.run(collect(REGISTRY["jobicy"]))[0].payload)
    assert item["title"] == "Closer" and item["company"] == "Acme" and item["work_mode"] == "remote"
    item = normalize("remotefirstjobs", asyncio.run(collect(REGISTRY["remotefirstjobs"]))[0].payload)
    assert item["company"] == "Acme" and item["title"] == "Closer"
    item = normalize("jobspresso", asyncio.run(collect(REGISTRY["jobspresso"]))[0].payload)
    assert item["title"] == "Closer de ventas" and item["company"] == "Empresa no indicada"


def test_ats_connectors_and_watchlist(monkeypatch, tmp_path):
    monkeypatch.setenv("CAREER_DB", str(tmp_path / "watchlist.sqlite3"))
    monkeypatch.setenv("CAREER_ORG_ID", "local-wdc")
    from career.db import initialize
    initialize()
    async def fake_fetch(url, params=None, ttl=900):
        if "boards-api.greenhouse.io" in url:
            return json.dumps({"jobs": [{"id": 7, "title": "Sales Closer", "company_name": "Acme",
                "absolute_url": "https://boards.greenhouse.io/acme/jobs/7", "location": {"name": "Spain (Remote)"},
                "first_published": "2026-09-22T10:00:00-04:00"}]})
        if "ashbyhq.com" in url:
            return json.dumps({"jobs": [{"id": "ab", "title": "Setter", "jobUrl": "https://jobs.ashbyhq.com/acme/ab",
                "location": "Madrid", "publishedAt": "2026-09-21T08:00:00.000+00:00", "workplaceType": "Hybrid"}]})
        if "recruitee.com" in url:
            return json.dumps({"offers": [{"id": 99, "title": "Account Executive",
                "careers_url": "https://jobs.acme.co/o/ae", "description": "<p>Ventas</p>",
                "locations": [{"name": "Anywhere in Europe"}], "remote": True,
                "published_at": "2026-05-28 13:33:37 UTC",
                "salary": {"min": 60000, "max": 80000, "currency": "EUR", "period": "year"}}]})
        if url == "https://careers.monestcapital.com/jobs.rss":
            return ('<?xml version="1.0" encoding="UTF-8"?><rss version="2.0" xmlns:tt="https://teamtailor.com/locations">'
                    '<channel><title>Monest Capital</title><item><title>Closer de ventas</title>'
                    '<link>https://careers.monestcapital.com/jobs/123</link><guid>abc-uuid-123</guid>'
                    '<description>&lt;p&gt;Inversor PRO. Venta consultiva.&lt;/p&gt;</description>'
                    '<pubDate>Tue, 23 Sep 2026 08:00:00 GMT</pubDate>'
                    '<tt:location><tt:name>Remoto</tt:name></tt:location></item></channel></rss>')
        raise AssertionError(url)
    monkeypatch.setattr("career.connectors.fetch", fake_fetch)
    async def collect(connector):
        return [row async for batch in connector.batches(SearchQuery()) for row in batch]
    gh = normalize("ats-greenhouse-acme", asyncio.run(collect(AtsCompany("greenhouse", "acme")))[0].payload)
    assert gh["company"] == "Acme" and gh["locations"] == ["Spain (Remote)"]
    ab = normalize("ats-ashby-acme", asyncio.run(collect(AtsCompany("ashby", "acme")))[0].payload)
    assert ab["work_mode"] == "Hybrid" and ab["source_url"].startswith("https://jobs.ashbyhq.com/")
    rc = normalize("ats-recruitee-acme", asyncio.run(collect(AtsCompany("recruitee", "acme")))[0].payload)
    assert rc["work_mode"] == "remote" and rc["salary_min"] == 60000 and rc["salary_currency"] == "EUR"
    # Dominio propio de Teamtailor (Monest Capital / Inversor PRO, hallazgo Manu).
    monest_conn = watchlist_connectors()["ats-teamtailor-monestcapital"]
    monest = normalize(monest_conn.code, asyncio.run(collect(monest_conn))[0].payload)
    assert monest["title"] == "Closer de ventas" and monest["company"] == "Monest Capital / Inversor PRO"
    assert monest["source_url"] == "https://careers.monestcapital.com/jobs/123" and monest["work_mode"] == "remote"
    assert set(watchlist_connectors()) == {"ats-greenhouse-stripe", "ats-greenhouse-typeform", "ats-greenhouse-wallapop",
                                           "ats-greenhouse-cabify", "ats-ashby-pennylane", "ats-recruitee-livestorm",
                                           "ats-teamtailor-producthackers", "ats-teamtailor-monestcapital"}


def test_provider_contracts(client, candidate, monkeypatch):
    async def fake_fetch(url, params=None, ttl=900):
        if "himalayas" in url:
            return json.dumps({"jobs": [candidate.payload]})
        if "workingnomads" in url:
            return json.dumps([{"url": "https://example.com/wn", "title": "Sales Closer", "company_name": "Closer Team",
                                "description": "Remote sales role", "location": "Anywhere", "pub_date": "2026-09-22T10:00:00-00:00"}])
        if "virtualvocations" in url:
            return '<rss><channel><item><title>Acme: Closer</title><link>https://example.com/vv</link><description>Remote sales role</description></item></channel></rss>'
        if "remote.jobs" in url:
            return '<rss><channel><item><title>Setter remoto</title><link>https://example.com/rj</link><description>Agenda demos para equipo comercial</description></item></channel></rss>'
        return '<rss><channel><item><title>Company: Sales Closer</title><link>https://example.com/job</link><region>Anywhere in the World</region><description>Sales position</description></item></channel></rss>'
    monkeypatch.setattr(connectors, "fetch", fake_fetch)
    async def collect(connector):
        return [row async for batch in connector.batches(SearchQuery()) for row in batch]
    himalayas = asyncio.run(collect(Himalayas()))
    assert len(himalayas) == 1
    wwr = asyncio.run(collect(WeWorkRemotely()))
    normalized = normalize("weworkremotely", wwr[0].payload)
    assert normalized["company"] == "Company"
    assert normalized["locations"] == ["Anywhere in the World"]
    # New boards requested by WDC: public feed verified 2026-09-23.
    wn = asyncio.run(collect(REGISTRY["workingnomads"]))
    assert len(wn) == 1
    wn_item = normalize("workingnomads", wn[0].payload)
    assert wn_item["company"] == "Closer Team"
    assert wn_item["published_at"] and wn_item["source_url"]
    for code in ["virtualvocations", "remotejobs"]:
        rows = asyncio.run(collect(REGISTRY[code]))
        assert len(rows) == 1, code
        item = normalize(code, rows[0].payload)
        assert item["source_url"].startswith("https://example.com/"), code
        assert item["work_mode"] == "remote", code
    vv = normalize("virtualvocations", {"title": "Acme: Closer", "link": "https://example.com/vv", "description": "Remote sales role"})
    assert vv["company"] == "Acme" and vv["title"] == "Closer"
    rj = normalize("remotejobs", {"title": "Setter remoto", "link": "https://example.com/rj", "description": "Agenda demos"})
    assert rj["company"] == "Empresa no indicada" and rj["title"] == "Setter remoto"


def test_spanish_bilingual_europe_and_work_mode(candidate):
    item = normalize('himalayas', candidate.payload)
    query = SearchQuery(role='sales', market='EU', require_spanish=True)
    item['locations'] = ['Berlin, Germany']
    item['title'] = 'Comercial digital'
    item['description'] = 'Buscamos experiencia en ventas para nuestro equipo. Trabajo híbrido.'
    item['work_mode'] = 'hybrid'
    assert match(item, query)
    assert match(item, query.model_copy(update={'work_mode': 'remote'})) is None
    item['description'] = 'English-speaking sales team. On-site only.'
    item['work_mode'] = 'onsite'
    assert match(item, query) is None
    item['description'] = 'Bilingual Spanish and English. Work from home.'
    item['work_mode'] = ''
    assert match(item, query)
    item['description'] = 'English only. Remote.'
    assert match(item, query) is None
    item['description'] = 'Remote. Spanish not required.'
    assert match(item, query) is None
    item['locations'] = ['United States']
    item['description'] = 'Spanish-speaking closer. Remote.'
    assert match(item, query) is None


def test_portuguese_is_not_spanish_and_sales_management_is_not_default(candidate):
    item = normalize('himalayas', candidate.payload)
    item['title'] = 'Executivo Comercial'
    item['description'] = 'Buscamos experiência com clientes e requisitos. Trabalho remoto.'
    assert match(item, SearchQuery(role='digital_sales', require_spanish=True)) is None
    item['title'] = 'Líder de Equipo de Ventas'
    item['description'] = 'Dirige nuestro equipo de ventas en remoto.'
    assert match(item, SearchQuery(role='digital_sales', require_spanish=True)) is None
    item['title'] = 'Coordinador de inteligencia comercial'
    assert match(item, SearchQuery(role='digital_sales', require_spanish=True)) is None


@pytest.mark.parametrize('title', ['Closer de Ventas', 'Closer digital', 'High Ticket Closer', 'Appoitment Setter', 'Appointment Setter', 'Setter', 'Comercial Digital', 'Cerrador', 'Cerradora', 'Venta de Alto Valor'])
def test_requested_aliases(candidate, title):
    item = normalize('himalayas', candidate.payload)
    item['title'] = title
    assert match(item, SearchQuery(role='sales', require_spanish=True))


def test_sources_secrets_isolation_and_availability(client, monkeypatch):
    body = {'name': 'InfoJobs prueba', 'source': 'infojobs', 'token': 'fake-unit-secret', 'actor_id': 'user/actor', 'dataset_id': 'unit-dataset', 'urls': ['https://www.infojobs.net/ofertas-trabajo/closer']}
    response = client.post('/v1/source-connections', json=body)
    assert response.status_code == 200, response.text
    value = response.json()
    assert value['has_token'] and 'fake-unit-secret' not in response.text and '"token"' not in response.text
    connection_id = value['id']
    with connect() as db:
        assert 'fake-unit-secret' not in db.execute('SELECT payload FROM source_connections').fetchone()[0]
    assert any(s['code'] == connection_id and s['available'] for s in client.get('/v1/config').json()['sources'])
    run = create(client, {'sources': [connection_id]}, 'configured-run-123')
    assert run['sources'][0]['name'] == 'InfoJobs prueba'
    from career.sources import secret_path
    assert secret_path('org-a', connection_id).stat().st_mode & 0o777 == 0o600
    monkeypatch.setenv('CAREER_ORG_ID', 'org-b')
    assert client.get('/v1/source-connections').json() == []
    assert client.patch('/v1/source-connections/' + connection_id, json=body).status_code == 404
    assert client.post('/v1/source-connections/' + connection_id + '/check', json={}).status_code == 404
    assert client.post('/v1/search-runs', json={'query': {'sources': [connection_id]}, 'idempotency_key': 'cross-org-test'}).status_code == 422


@pytest.mark.parametrize('patch', [
    {'urls': ['http://example.com']}, {'urls': ['https://127.0.0.1/private']},
    {'urls': ['https://example.com?token=secret']}, {'actor_input': {'cookies': 'secret'}},
    {'mapping': {'source_url': 'url', 'bogus': 'password'}}, {'max_items': 501},
])
def test_connection_invalid_input(client, patch):
    assert client.post('/v1/source-connections', json={'name': 'Bad', **patch}).status_code == 422


def test_apify_launch_check_import_and_no_paid_search(client, monkeypatch):
    from career import sources
    calls = []
    async def fake_request(method, path, token, **kwargs):
        calls.append((method, path, kwargs))
        assert token == 'fake-unit-secret'
        if method == 'POST':
            assert kwargs['payload']['startUrls'] == [{'url': 'https://example.com/jobs'}]
            assert kwargs['params']['maxTotalChargeUsd'] == 1
            return {'data': {'id': 'run-test', 'status': 'RUNNING'}}
        if path.startswith('actor-runs'):
            return {'data': {'id': 'run-test', 'status': 'SUCCEEDED', 'defaultDatasetId': 'dataset-test', 'usageTotalUsd': .01}}
        return [{'title': 'Closer digital', 'companyName': 'Test', 'description': 'Buscamos experiencia en ventas para nuestro equipo. Trabajo remoto con clientes.', 'url': 'https://example.com/job', 'location': 'Spain', 'postedAt': now(), 'workMode': 'remote'}]
    monkeypatch.setattr(sources, 'apify_request', fake_request)
    body = {'name': 'Apify fixture', 'token': 'fake-unit-secret', 'actor_id': 'user/actor', 'urls': ['https://example.com/jobs']}
    value = client.post('/v1/source-connections', json=body).json()
    path = '/v1/source-connections/' + value['id']
    assert client.post(path + '/scrape').status_code == 200
    assert client.post(path + '/scrape').status_code == 409
    assert client.patch(path, json=body).status_code == 409
    check = client.post(path + '/check')
    assert check.status_code == 200, check.text
    assert check.json()['connection']['dataset_id'] == 'dataset-test'
    run = create(client, {'sources': [value['id']], 'require_spanish': True})
    asyncio.run(execute(run['id'], 'org-a', SearchQuery(**run['query']), sources.registry('org-a')))
    result = load_run(run['id'], 'org-a')
    assert result['status'] == 'completed' and len(result['jobs']) == 1
    assert len([c for c in calls if c[0] == 'POST']) == 1  # searching only reads existing data
    assert client.patch(path, json={**body, 'token': '', 'dataset_id': 'dataset-test'}).json()['has_token']
    assert not client.patch(path, json={**body, 'token': '', 'clear_token': True}).json()['has_token']


def test_uncertain_scrape_is_not_retried(client, monkeypatch):
    from career import sources
    async def failed(*args, **kwargs):
        raise ProviderError('timeout')
    monkeypatch.setattr(sources, 'apify_request', failed)
    value = client.post('/v1/source-connections', json={'name': 'Uncertain', 'actor_id': 'user/actor', 'token': 'fake-unit-secret'}).json()
    path = '/v1/source-connections/' + value['id']
    assert client.post(path + '/scrape').status_code == 502
    assert client.post(path + '/scrape').status_code == 409
    assert client.post(path + '/resolve', json={'reviewed_apify_console': False}).status_code == 409
    assert client.post(path + '/resolve', json={'reviewed_apify_console': True}).status_code == 200


def test_new_public_connector_contracts(client, monkeypatch):
    from career.connectors import PublicJSON, GetOnBoard
    async def fake_fetch(url, params=None, ttl=900):
        if 'getonbrd' in url:
            return json.dumps({'data': [{'id': 'closer', 'attributes': {'title': 'Closer digital', 'description': 'Buscamos experiencia en ventas para el equipo.', 'remote_modality': 'hybrid', 'countries': ['Chile'], 'published_at': 1789684387}, 'links': {'public_url': 'https://www.getonbrd.com/jobs/test'}}], 'meta': {'total_pages': 1}})
        return json.dumps({'jobs': [{'id': 1, 'title': 'Closer', 'url': 'https://remotive.com/job/1', 'description': 'Spanish', 'candidate_required_location': 'Spain', 'publication_date': now()}]})
    monkeypatch.setattr(connectors, 'fetch', fake_fetch)
    async def collect(c):
        return [row async for batch in c.batches(SearchQuery()) for row in batch]
    rows = asyncio.run(collect(GetOnBoard()))
    item = normalize('getonbrd-public', rows[0].payload)
    assert item['work_mode'] == 'hybrid' and item['company'] == 'Empresa no indicada'
    rows = asyncio.run(collect(PublicJSON('remotive', 'Remotive', 'https://remotive.com/api/remote-jobs', 'jobs')))
    assert normalize('remotive', rows[0].payload)['work_mode'] == 'remote'


def test_jobtoday_embedded_listing_contract(client, monkeypatch):
    """Contrato del listado embebido de JobToday verificado en vivo 2026-09-23."""
    from career.connectors import JobToday
    payload = {'key': 'DLpjnX', 'role': 'Closer de ventas', 'companyName': 'Hispaconex',
               'description': 'Buscamos closer de ventas remoto. Experiencia comercial.',
               'canonicalUrl': '/es/trabajo/closer-hispaconex-DLpjnX',
               'address': 'Tetuán, Madrid, Comunidad de Madrid, Spain',
               'createDate': 1788358516586, 'postedSecondsAgo': 1801361,
               'salary': {'from': 1500, 'to': 1500, 'currencyCode': 'EUR', 'period': 'MONTHLY', 'isValid': True}}
    html = '<html><script id="__NEXT_DATA__" type="application/json">' + json.dumps(
        {'props': {'pageProps': {'feed': {'sections': [{'items': [{'payload': payload}]}]}}}}) + '</script></html>'
    async def fake_fetch(url, params=None, ttl=900, user_agent=None):
        assert 'jobtoday.com/es/trabajos-closer-venta' in url
        return html
    monkeypatch.setattr(connectors, 'fetch', fake_fetch)
    connector = JobToday('jobtoday', 'JobToday España', 'closer-venta')
    async def collect():
        return [row async for batch in connector.batches(SearchQuery()) for row in batch]
    rows = asyncio.run(collect())
    assert len(rows) == 1 and rows[0].external_id == 'DLpjnX'
    item = normalize('jobtoday', rows[0].payload)
    assert item['title'] == 'Closer de ventas' and item['company'] == 'Hispaconex'
    assert item['source_url'].startswith('https://jobtoday.com/es/trabajo/')
    assert item['salary_min'] == 1500 and item['salary_currency'] == 'EUR' and item['salary_period'] == 'monthly'
    assert item['locations'] == ['Tetuán, Madrid, Comunidad de Madrid, Spain']
    assert item['published_at']  # de createDate o inferida de postedSecondsAgo
    # Sin salary válido no se inventa compensación; sin address no hay ubicación fabricada.
    payload2 = dict(payload, salary={'from': 0, 'to': 0, 'isValid': False}, address=None, createDate=None, postedSecondsAgo=None)
    html2 = html.replace(json.dumps(payload), json.dumps(payload2))
    monkeypatch.setattr(connectors, 'fetch', lambda url, params=None, ttl=900, user_agent=None: _async_return(html2))
    rows = asyncio.run(collect())
    item2 = normalize('jobtoday', rows[0].payload)
    assert item2['salary_min'] is None and item2['locations'] == []


async def _drain(connector, query=None):
    return [row async for batch in connector.batches(query or SearchQuery()) for row in batch]


def _async_return(value):
    async def inner(*args, **kwargs):
        return value
    return inner()


SERPAPI_JOBS_PAYLOAD = json.dumps({
    "search_metadata": {"id": "abc123"},
    "jobs_results": [
        {"title": "Closer de ventas", "via": "Academia Ejemplo",
         "description": "<p>Remoto. Venta de formacion online con leads cualificados.</p>",
         "job_id": "Vk123", "link": "https://www.linkedin.com/jobs/view/vk123",
         "location": "Espa\u00f1a Remoto",
         "detected_extensions": {"posted_at": "2026-09-23T08:00:00Z"}},
        {"title": "Asesor de Admisiones", "via": "Bootcamp Ejemplo",
         "description": "<p>Captacion de alumnos remoto.</p>",
         "job_id": "Ab456", "apply_options": [{"link": "https://bootcamp.example/apply"}],
         "location": "Madrid, Espa\u00f1a",
         "detected_extensions": {"posted_at_relative": "hace 2 dias"}},
    ],
})


def test_serpapi_connector_maps_google_jobs(client, monkeypatch, tmp_path):
    """El conector google consume engine=google_jobs y mapea jobs_results sin inventar datos."""
    monkeypatch.setenv("CAREER_ORG_ID", "local-wdc")
    calls = []
    async def fake_fetch(url, params=None, ttl=900, user_agent=None):
        calls.append((url, dict(params or {})))
        assert 'serpapi.com/search.json' in url
        assert params['engine'] == 'google_jobs'
        assert params['api_key'] == 'test-key'
        return SERPAPI_JOBS_PAYLOAD
    monkeypatch.setattr(connectors, 'fetch', fake_fetch)
    from career.integrations import save_secret
    save_secret('serpapi', SecretStr('test-key'), False, 'local-wdc')
    from career.connectors import google_serpapi_connector
    connector = google_serpapi_connector('site:linkedin.com/jobs/view "closer de ventas" Espa\u00f1a')
    async def collect():
        return [row async for batch in connector.batches(SearchQuery()) for row in batch]
    rows = asyncio.run(collect())
    assert len(rows) == 2
    item = normalize('google', rows[0].payload)
    assert item['title'] == 'Closer de ventas' and item['company'] == 'Academia Ejemplo'
    assert item['source_url'] == 'https://www.linkedin.com/jobs/view/vk123'
    assert item['published_at'] and item['work_mode'] == 'remote'
    item2 = normalize('google', rows[1].payload)
    assert item2['source_url'] == 'https://bootcamp.example/apply', 'apply_options como enlace alternativo'
    assert item2['published_at'], 'fecha relativa explícita se convierte'
    # Un dork = una búsqueda: sin token la fuente no lanza nada.
    save_secret('serpapi', SecretStr(''), True, 'local-wdc')
    with pytest.raises(ProviderError):
        asyncio.run(collect())


def test_google_run_respects_budget_and_registry(client, monkeypatch):
    """El endpoint google/run reserva presupuesto y el registry expone el conector con clave."""
    monkeypatch.setenv("CAREER_ORG_ID", "local-wdc")
    from career.integrations import save_secret
    from career.sources import registry
    from career.connectors import google_serpapi_connector
    save_secret('serpapi', SecretStr('test-key'), False, 'local-wdc')
    reg = registry('local-wdc')
    assert 'google' in reg, 'con clave guardada, google entra en el registry'
    assert reg['google'].provider == 'serpapi_google_jobs_v1'
    assert len(reg['google'].dorks) >= 12
    # Sin clave el conector se niega: nunca lanza busquedas de pago sin credencial.
    from career.sources import secret_path
    secret_path('local-wdc', 'account-serpapi').unlink()
    connector = google_serpapi_connector('site:linkedin.com/jobs/view "closer" Espana')
    with pytest.raises(ProviderError, match='SerpApi'):
        asyncio.run(_drain(connector))


def test_google_daily_collection_includes_serpapi_with_key(client, monkeypatch):
    """La recogida diaria solo añade google si hay clave; nunca lanza sin credencial."""
    monkeypatch.setenv("CAREER_ORG_ID", "local-wdc")
    from career.collection import enqueue
    from career.db import connect as db_connect
    # Sin clave: no debe aparecer google.
    enqueue('local-wdc')
    with db_connect() as db:
        row = db.execute("SELECT sources FROM collection_runs WHERE org_id='local-wdc' ORDER BY created_at DESC LIMIT 1").fetchone()
        sources = json.loads(row['sources'])
    assert all(s['code'] != 'google' for s in sources)
    # Con clave y conexion activa: google entra en la lista (INSERT OR IGNORE reutiliza el run del día,
    # así que borramos para forzar re-encolado con la nueva lista de fuentes).
    from career.integrations import save_secret
    from career.sources import put_connection, get_or_create_google_connection
    save_secret('serpapi', SecretStr('test-key'), False, 'local-wdc')
    value = get_or_create_google_connection('local-wdc')
    value['enabled'] = True
    put_connection(value, 'local-wdc')
    with db_connect() as db:
        db.execute("DELETE FROM collection_runs WHERE org_id='local-wdc'")
    enqueue('local-wdc')
    with db_connect() as db:
        row = db.execute("SELECT sources FROM collection_runs WHERE org_id='local-wdc' ORDER BY created_at DESC LIMIT 1").fetchone()
        sources2 = json.loads(row['sources'])
    assert any(s['code'] == 'google' for s in sources2)


COMPANIES_PAYLOAD = json.dumps({
    "search_metadata": {"id": "co1"},
    "jobs_results": [
        {"title": "Closer senior", "via": "Portal Ejemplo",
         "description": "<p>Remoto. Venta consultiva con leads propios.</p>",
         "job_id": "Co789", "link": "https://www.linkedin.com/jobs/view/co789",
         "location": "España (Remoto)",
         "detected_extensions": {"posted_at_relative": "hace 3 dias"}},
    ],
})


def _add_linkedin_guest_company(slug, name, priority=1):
    """Empresa objetivo de vigilancia LinkedIn para tests (datos neutros, sin tenant real)."""
    with connect() as db:
        db.execute("INSERT OR IGNORE INTO company_target (slug,name,website,sector,country,ticket,careers_url,priority,score,monitoring,enabled,created_at) "
                   "VALUES (?,?,?,?,?,?,?,?,?,?,1,strftime('%Y-%m-%dT%H:%M:%f+00:00','now'))",
                   (slug, name, '', 'formacion', 'España', '', '', priority, 0, 'linkedin_guest'))


def test_google_company_dorks_anchor_exact_company(client, monkeypatch):
    """Patrón del usuario: "empresa" (rol OR ...) España, en orden estable por prioridad."""
    monkeypatch.setenv("CAREER_ORG_ID", "local-wdc")
    from career.sources import google_company_dorks, company_watch_targets
    pairs = google_company_dorks([('Academia Ejemplo', 'Academia Ejemplo'), ('Bootcamp Dos', 'Bootcamp Dos')])
    assert [p[0] for p in pairs] == ['Academia Ejemplo', 'Bootcamp Dos']
    dork = pairs[0][2]
    assert dork.startswith('"Academia Ejemplo" (') and dork.endswith(') España'), dork
    roles = dork[len('"Academia Ejemplo" ('):-len(') España')].split(' OR ')
    assert sorted(roles) == ['"asesor comercial"', '"closer"', '"sales advisor"', '"setter"'], dork
    # Las empresas salen de company_target con canal linkedin_guest y enabled.
    _add_linkedin_guest_company('academia-ejemplo', 'Academia Ejemplo', priority=1)
    _add_linkedin_guest_company('bootcamp-dos', 'Bootcamp Dos', priority=2)
    with connect() as db:
        db.execute("UPDATE company_target SET enabled=0 WHERE slug='bootcamp-dos'")
    targets = company_watch_targets('local-wdc')
    assert ('Academia Ejemplo', 'Academia Ejemplo') in targets
    assert all(name != 'Bootcamp Dos' for name, _ in targets), 'deshabilitada no se vigila'


def test_google_companies_connector_one_search_per_company(client, monkeypatch):
    """Una búsqueda google_jobs por empresa, atribuida a la empresa objetivo; sin clave no lanza nada."""
    monkeypatch.setenv("CAREER_ORG_ID", "local-wdc")
    calls = []
    async def fake_fetch(url, params=None, ttl=900, user_agent=None):
        calls.append((url, dict(params or {})))
        assert 'serpapi.com/search.json' in url
        assert params['engine'] == 'google_jobs'
        assert params['api_key'] == 'test-key'
        prefix = params['q'].split(' (')[0]
        assert prefix in ('"Academia Ejemplo"', '"Bootcamp Dos"'), params['q']
        assert prefix.count('"') == 2 and params['q'].endswith('España'), 'empresa anclada como frase exacta'
        await asyncio.sleep(0)  # El rate-limit real (2 s) no aplica en tests.
        return COMPANIES_PAYLOAD
    monkeypatch.setattr(connectors, 'fetch', fake_fetch)
    from career.integrations import save_secret
    save_secret('serpapi', SecretStr('test-key'), False, 'local-wdc')
    from career.connectors import GoogleCompaniesSerpApi
    connector = GoogleCompaniesSerpApi([('Academia Ejemplo', 'Academia Ejemplo'), ('Bootcamp Dos', 'Bootcamp Dos')])
    rows = asyncio.run(_drain(connector))
    assert len(calls) == 2, 'una búsqueda por empresa objetivo'
    assert {c[1]['q'].split(' (')[0] for c in calls} == {'"Academia Ejemplo"', '"Bootcamp Dos"'}
    assert len(rows) == 2
    for row in rows:
        assert row.payload['company_name'] == row.payload['_search_company'], 'atribución a la empresa buscada'
        item = normalize('google-companies', row.payload)
        assert item['company'] == row.payload['_search_company']
        assert item['title'] == 'Closer senior'
        assert item['source_url'] == 'https://www.linkedin.com/jobs/view/co789'
        assert item['published_at'], 'fecha relativa se convierte'
        assert item['work_mode'] == 'remote'
    assert {r.payload['_search_company'] for r in rows} == {'Academia Ejemplo', 'Bootcamp Dos'}
    # Sin clave: la vigilancia no lanza búsquedas de pago.
    save_secret('serpapi', SecretStr(''), True, 'local-wdc')
    with pytest.raises(ProviderError):
        asyncio.run(_drain(connector))


def test_google_companies_daily_collection_gated_by_key_and_enabled(client, monkeypatch):
    """La recogida diaria solo incluye google-companies con clave y conexión activa."""
    monkeypatch.setenv("CAREER_ORG_ID", "local-wdc")
    from career.collection import enqueue
    from career.db import connect as db_connect
    _add_linkedin_guest_company('academia-ejemplo', 'Academia Ejemplo')
    enqueue('local-wdc')
    with db_connect() as db:
        row = db.execute("SELECT sources FROM collection_runs WHERE org_id='local-wdc' ORDER BY created_at DESC LIMIT 1").fetchone()
        sources = json.loads(row['sources'])
    assert all(s['code'] != 'google-companies' for s in sources), 'sin clave no entra'
    from career.integrations import save_secret
    save_secret('serpapi', SecretStr('test-key'), False, 'local-wdc')
    with db_connect() as db:
        db.execute("DELETE FROM collection_runs WHERE org_id='local-wdc'")
    enqueue('local-wdc')
    with db_connect() as db:
        row = db.execute("SELECT sources FROM collection_runs WHERE org_id='local-wdc' ORDER BY created_at DESC LIMIT 1").fetchone()
        sources2 = json.loads(row['sources'])
    assert any(s['code'] == 'google-companies' for s in sources2)


def test_launch_serpapi_companies_reserves_and_closes_on_run_completion(client, monkeypatch):
    """El lanzamiento manual reserva presupuesto y collect cierra la reserva con coste real."""
    monkeypatch.setenv("CAREER_ORG_ID", "local-wdc")
    from career.integrations import save_secret
    save_secret('serpapi', SecretStr('test-key'), False, 'local-wdc')
    _add_linkedin_guest_company('academia-ejemplo', 'Academia Ejemplo', priority=1)
    _add_linkedin_guest_company('bootcamp-dos', 'Bootcamp Dos', priority=2)
    _add_linkedin_guest_company('colegio-tres', 'Colegio Tres', priority=3)
    async def fake_rate():
        return 1.10
    monkeypatch.setattr('career.collection.euro_rate', fake_rate)
    from career.sources import launch_serpapi_companies
    result = asyncio.run(launch_serpapi_companies('local-wdc'))
    assert result['pending_companies'] == 0, 'con 1 EUR/día caben las 3 búsquedas'
    assert 'Vigilancia de 3 empresas' in result['message']
    with connect() as db:
        reservation = db.execute("SELECT id,status,cap_usd FROM scrape_budget WHERE org_id='local-wdc' AND connection_id='google-companies-serpapi'").fetchone()
        assert reservation['status'] == 'RUNNING'
        assert reservation['cap_usd'] + 1e-9 >= 3 * 0.025, 'cap cubre una búsqueda por empresa'
    # El run encolado lo consume el worker (re-ranking del catálogo, camino sin registry);
    # al terminalizar el run, execute() cierra la reserva SerpApi con el cap reservado.
    asyncio.run(execute(result['run_id'], 'local-wdc', SearchQuery()))
    with connect() as db:
        closed = db.execute("SELECT status,actual_usd FROM scrape_budget WHERE org_id='local-wdc' AND id=?", (reservation['id'],)).fetchone()
    assert closed['status'] == 'SUCCEEDED', closed['status']
    assert closed['actual_usd'] and closed['actual_usd'] > 0
    # El run_status de la conexión interna vuelve a terminal: sin esto, el segundo intento
    # daría 409 "ejecución pendiente" para siempre (también afectaba a los dorks google).
    from career.sources import get_connection, check_connection
    conn = get_connection('google-companies-serpapi', 'local-wdc')
    assert conn['run_status'] == 'SUCCEEDED', conn['run_status']
    check = asyncio.run(check_connection('google-companies-serpapi', 'local-wdc'))
    assert 'SUCCEEDED' in check['message']
    # La red de seguridad en collect() es idempotente: nada ya RUNNING queda sin cerrar.
    from career.collection import enqueue, collect
    with connect() as db:
        db.execute("DELETE FROM collection_runs WHERE org_id='local-wdc'")
    enqueue('local-wdc')
    with connect() as db:
        collection_run = db.execute("SELECT id FROM collection_runs WHERE org_id='local-wdc' ORDER BY created_at DESC LIMIT 1").fetchone()['id']
    asyncio.run(collect(collection_run, 'local-wdc'))
    with connect() as db:
        closed2 = db.execute("SELECT status,actual_usd FROM scrape_budget WHERE org_id='local-wdc' AND id=?", (reservation['id'],)).fetchone()
    assert closed2['status'] == 'SUCCEEDED' and closed2['actual_usd'] == closed['actual_usd']
